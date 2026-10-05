"""Genera dos Word: resumen con viñetas y resumen narrativo."""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_PUNTOS = Path(__file__).resolve().parent / "Resumen-Proyecto-Fichajes.docx"
OUT_NARRATIVO = Path(__file__).resolve().parent / "Resumen-Proyecto-Fichajes-narrativo.docx"


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def portada(doc):
    title = doc.add_heading("Sistema de fichajes (control horario)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, "Autora: Svetlana", bold=True)
    add_para(doc, "Fecha: 1 de octubre de 2026")
    add_para(doc, "Stack: PostgreSQL + Node.js (Express, TypeScript) + Frontend HTML/JS")
    doc.add_paragraph()


def build_doc_puntos():
    doc = Document()
    portada(doc)

    doc.add_heading("1. Objetivo del proyecto", 1)
    add_para(
        doc,
        "Aplicación de fichajes (entrada/salida) con PostgreSQL, API REST con JWT e interfaz web. "
        "Prioridad: parte del empleado; administración y RRHH en fase posterior.",
    )

    doc.add_heading("2. Diseño de la base de datos", 1)
    doc.add_heading("2.1 Tablas principales (backend/init.sql)", 2)
    add_bullets(
        doc,
        [
            "usuarios: datos del trabajador, rol, centro, departamento, password_hash.",
            "fichajes: registros inmutables ENTRADA/SALIDA con hora del servidor.",
            "solicitudes_modificacion: correcciones cuando se olvidó fichar (RRHH, fase futura).",
        ],
    )
    doc.add_heading("2.2 Tipos ENUM", 2)
    add_bullets(
        doc,
        [
            "tipo_fichaje: ENTRADA, SALIDA",
            "rol_usuario: EMPLEADO, ADMIN, RRHH",
            "estado_solicitud: PENDIENTE, APROBADA, RECHAZADA",
        ],
    )
    doc.add_heading("2.3 Datos de ejemplo (backend/seed.sql)", 2)
    add_bullets(
        doc,
        [
            "5 usuarios, 5 fichajes, 5 solicitudes.",
            "Contraseña de prueba (todos): Demo1234",
            "Ejemplo empleado: ana.martinez@empresa.com / Demo1234",
            "Recarga: npm run db:seed (desde backend)",
        ],
    )

    doc.add_heading("3. Infraestructura Docker", 1)
    add_bullets(
        doc,
        [
            "docker-compose.yml: PostgreSQL fichajes_postgres, puerto 5432, base fichajes_db.",
            "Init: 01-init.sql (esquema) y 02-seed.sql (datos) al crear el volumen.",
            "DATABASE_URL local: postgresql://postgres:mi_password_secreto@localhost:5432/fichajes_db",
        ],
    )

    doc.add_heading("4. Backend (API Node.js + TypeScript)", 1)
    doc.add_heading("4.1 Arranque", 2)
    add_bullets(
        doc,
        [
            "src/index.ts: dotenv y listen.",
            "src/app.ts: Express, CORS, rutas.",
            "Comando: npm run dev (tsx + watch, puerto 3000).",
        ],
    )
    doc.add_heading("4.2 Capas", 2)
    add_bullets(
        doc,
        [
            "Routes: auth.routes.ts, fichajes.routes.ts",
            "Services: auth.service.ts, fichaje.service.ts",
            "Repositories: usuario.repository.ts, fichaje.repository.ts",
            "Config: database.ts, env.ts",
        ],
    )
    doc.add_heading("4.3 Autenticación JWT", 2)
    add_bullets(
        doc,
        [
            "POST /api/auth/login → token + user (sin password_hash).",
            "GET /api/auth/me → requiere Authorization: Bearer <token>.",
            "auth.middleware.ts valida JWT y rellena req.user.",
            "Contraseñas con bcrypt frente a password_hash.",
        ],
    )
    doc.add_heading("4.4 Fichajes", 2)
    add_bullets(
        doc,
        [
            "POST /api/fichajes — registrar ENTRADA/SALIDA (JWT).",
            "GET /api/fichajes/ultimo — último fichaje del usuario.",
            "GET /api/fichajes/mios — historial.",
            "Regla: no dos ENTRADA ni dos SALIDA seguidas (409).",
        ],
    )
    add_para(doc, "GET /api/health → { ok: true }")

    doc.add_heading("5. Frontend (carpeta frontend)", 1)
    add_bullets(
        doc,
        [
            "serve.py → http://localhost:8090 (no abrir HTML con file://).",
            "login.html + login.js: acceso contra la API.",
            "index.html + app.js: pantalla de fichaje.",
            "auth.js: API_BASE, loginWithApi, apiFetch, token en sessionStorage.",
            "Funciones: fichar, obtenerUltimoFichaje, obtenerMisFichajes.",
            "server-popup.js: alert con respuestas del servidor.",
            "Protección index: isUserLoggedIn() por token; sin sesión → login.html.",
        ],
    )
    add_para(
        doc,
        "Pendiente: conectar botones Entrada/Salida e historial en app.js a la API "
        "(parte del historial sigue en localStorage).",
    )

    doc.add_heading("6. Cómo ejecutar en desarrollo", 1)
    add_bullets(
        doc,
        [
            "backend: docker compose up -d db",
            "backend: npm run dev (puerto 3000)",
            "frontend: python serve.py (puerto 8090)",
            "Navegador: http://localhost:8090/login.html",
        ],
    )

    doc.add_heading("7. Usuarios de prueba", 1)
    add_para(doc, "Contraseña: Demo1234", bold=True)
    add_bullets(
        doc,
        [
            "laura.garcia@empresa.com — ADMIN",
            "miguel.santos@empresa.com — RRHH",
            "ana.martinez@empresa.com — EMPLEADO",
            "carlos.fernandez@empresa.com — EMPLEADO",
            "elena.ruiz@empresa.com — EMPLEADO",
        ],
    )

    doc.add_heading("8. Problemas resueltos", 1)
    add_bullets(
        doc,
        [
            "Carpeta src/index.ts en lugar de archivo → corregido.",
            "init.sql duplicado en docker/ → no usado; activo backend/init.sql.",
            "401 login → contraseña seed unificada en Demo1234.",
            "localhost rechazó conexión → servir front en 8090 con serve.py.",
            "Token no proporcionado → Bearer en rutas protegidas tras login.",
        ],
    )

    doc.add_heading("9. Fases de trabajo", 1)
    add_bullets(
        doc,
        [
            "BD + seed + Docker",
            "Express + health",
            "Usuarios, login JWT, /me",
            "API fichajes",
            "Front: login real, popups, auth.js, protección index",
            "Próximo: fichaje e historial desde API en app.js",
        ],
    )

    doc.add_heading("10. Conclusiones", 1)
    add_para(
        doc,
        "Base funcional: BD normalizada, API con JWT y reglas de fichaje, web con login real. "
        "Siguiente paso: integrar botones de fichaje con el backend.",
    )
    return doc


def build_doc_narrativo():
    doc = Document()
    portada(doc)

    doc.add_heading("Introducción y objetivo", 1)
    add_para(
        doc,
        "Durante la práctica he desarrollado una aplicación de control horario para que los empleados "
        "registren su entrada y salida de la jornada. La idea es acercarme a un sistema serio a nivel "
        "de datos y seguridad, no limitarme a una página que guarde horas solo en el navegador. Por eso "
        "he separado el trabajo en tres bloques: base de datos, servidor (API) e interfaz web, que es "
        "la organización habitual en proyectos profesionales.",
    )
    add_para(
        doc,
        "He priorizado la experiencia del empleado: poder identificarse, fichar y consultar su propio "
        "historial. Las funciones de administración y de Recursos Humanos están previstas en el modelo "
        "de datos, pero las he dejado para una fase posterior.",
    )

    doc.add_heading("Diseño de la base de datos", 1)
    add_para(
        doc,
        "Empecé diseñando la base de datos en PostgreSQL con tres tablas principales: usuarios, "
        "fichajes y solicitudes de modificación. La contraseña se guarda como hash con bcrypt. Los "
        "fichajes son eventos con hora del servidor. El seed incluye cinco usuarios de prueba con "
        "contraseña Demo1234, por ejemplo ana.martinez@empresa.com.",
    )

    doc.add_heading("Infraestructura con Docker", 1)
    add_para(
        doc,
        "PostgreSQL corre en Docker (fichajes_db, puerto 5432). Al crear el volumen se aplican "
        "init.sql y seed.sql. El backend se conecta con DATABASE_URL en el archivo .env.",
    )

    doc.add_heading("Backend: API REST en Node.js", 1)
    add_para(
        doc,
        "Servidor con TypeScript y Express, organizado en rutas, servicios y repositorios. "
        "Autenticación JWT tras login; rutas protegidas con middleware Bearer. API de fichajes con "
        "regla de alternancia ENTRADA/SALIDA (error 409 si se repite la misma acción).",
    )

    doc.add_heading("Frontend: interfaz web del empleado", 1)
    add_para(
        doc,
        "Web servida con serve.py en el puerto 8090. Login real contra la API; token en "
        "sessionStorage. Popups con respuestas del servidor. Pantalla de fichaje protegida sin "
        "token. Pendiente: conectar botones e historial en app.js a la API.",
    )

    doc.add_heading("Cómo ejecuto el sistema en local", 1)
    add_para(
        doc,
        "docker compose up -d db, npm run dev en backend, python serve.py en frontend, y abro "
        "http://localhost:8090/login.html en el navegador.",
    )

    doc.add_heading("Usuarios de prueba", 1)
    add_para(
        doc,
        "Contraseña Demo1234 para todos. Perfiles ADMIN, RRHH y EMPLEADO en distintos centros, "
        "definidos en seed.sql.",
    )

    doc.add_heading("Dificultades y aprendizajes", 1)
    add_para(
        doc,
        "Corrección de index.ts como carpeta, unificación de init.sql, contraseña del seed, "
        "servidor front en 8090, y uso del header Authorization Bearer en rutas protegidas.",
    )

    doc.add_heading("Estado actual y conclusiones", 1)
    add_para(
        doc,
        "Base sólida: BD, API con JWT, login real en la web. Próximo paso: fichaje e historial "
        "desde app.js; después, flujo RRHH.",
    )
    return doc


def save_doc(doc, path: Path):
    try:
        doc.save(path)
        print(f"Generado: {path}")
    except PermissionError:
        print(f"No se pudo guardar (¿abierto en Word?): {path}")


def main():
    save_doc(build_doc_puntos(), OUT_PUNTOS)
    save_doc(build_doc_narrativo(), OUT_NARRATIVO)


if __name__ == "__main__":
    main()
